"""绘制 MAIZSIM Morris 敏感性分析论文图."""

from pathlib import Path
import argparse
import math

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from matplotlib.font_manager import FontProperties
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
from openpyxl.styles import Alignment, Font, PatternFill


ROOT_DIR = Path(__file__).resolve().parent
DEFAULT_ANALYSIS_DIR = (
    ROOT_DIR.parent
    / "Example input"
    / "SingleLayerLoam2D"
    / "sensitivity_runs"
    / "formal_r300_workers16_timeout300"
)
FONT_DIR = Path("C:/Windows/Fonts")
SIMSUN_CANDIDATES = [FONT_DIR / "simsun.ttc", FONT_DIR / "simsun.ttf"]
TIMES_CANDIDATES = [FONT_DIR / "times.ttf", FONT_DIR / "timesbd.ttf"]
METRIC_ORDER = [
    "max_lai",
    "final_shoot_dm",
    "final_ear_dm",
    "mean_top20_theta",
    "seasonal_transpiration_deficit",
    "seasonal_potential_transpiration",
    "seasonal_actual_transpiration",
]
METRIC_LABELS = {
    "max_lai": "Maximum LAI",
    "final_shoot_dm": "Final shoot dry matter",
    "final_ear_dm": "Final ear dry matter",
    "mean_top20_theta": "Mean 0-20 cm water content",
    "seasonal_transpiration_deficit": "Seasonal transpiration deficit",
    "seasonal_potential_transpiration": "Seasonal potential transpiration",
    "seasonal_actual_transpiration": "Seasonal actual transpiration",
}
PARAMETER_ORDER = [
    "JuvenileLeaves",
    "Rmax_LTAR",
    "Rmax_LTIR",
    "PhyllFrmTassel",
    "StayGreen",
    "LM_min",
    "Diffx",
    "Diffz",
    "thetaR",
    "thetaS",
    "Alfa",
    "n",
    "Ks",
]
PARAMETER_CATEGORY = {
    "JuvenileLeaves": "Crop phenology and leaf",
    "Rmax_LTAR": "Crop phenology and leaf",
    "Rmax_LTIR": "Crop phenology and leaf",
    "PhyllFrmTassel": "Crop phenology and leaf",
    "StayGreen": "Crop phenology and leaf",
    "LM_min": "Crop phenology and leaf",
    "Diffx": "Root distribution",
    "Diffz": "Root distribution",
    "thetaR": "Soil hydraulic",
    "thetaS": "Soil hydraulic",
    "Alfa": "Soil hydraulic",
    "n": "Soil hydraulic",
    "Ks": "Soil hydraulic",
}
CATEGORY_ORDER = ["Crop phenology and leaf", "Root distribution", "Soil hydraulic"]
PANEL_LABELS = ["a", "b", "c", "d", "e", "f", "g"]
TICK_SIZE = 7
LABEL_SIZE = 8.5
TITLE_SIZE = 8
LEGEND_SIZE = 7
ANNOTATION_SIZE = 6.5
BAR_EDGE_COLOR = "#333333"
GRID_COLOR = "#e6e6e6"


def parse_arguments():
    """解析命令行参数, 输出分析目录和绘图选项."""
    parser = argparse.ArgumentParser(description="Plot filtered Morris sensitivity figures.")
    parser.add_argument("--analysis-dir", default=str(DEFAULT_ANALYSIS_DIR), help="Directory containing Morris index CSV files.")
    parser.add_argument("--output-dir", default=None, help="Directory for output figures.")
    parser.add_argument(
        "--filtered-file",
        default="morris_indices_without_slow_trajectories.csv",
        help="Morris indices file after dropping slow trajectories.",
    )
    parser.add_argument(
        "--xlsx-file",
        default="morris_sensitivity_without_slow_tables.xlsx",
        help="Supplementary Excel workbook name or path.",
    )
    parser.add_argument("--top-n", default=6, type=int, help="Number of top parameters per metric in ranking figures.")
    return parser.parse_args()


def resolve_path(value, base_dir):
    """解析路径, 输入可以是绝对路径或相对路径."""
    path = Path(value)
    if path.is_absolute():
        return path
    return (base_dir / path).resolve()


def find_font(candidates, label):
    """检查字体文件是否存在, 缺失时直接报错."""
    for candidate in candidates:
        if candidate.is_file():
            return FontProperties(fname=str(candidate))
    names = ", ".join(str(candidate) for candidate in candidates)
    raise RuntimeError(f"{label} font missing: {names}")


def font_with_size(font, size):
    """复制字体并设置字号, 输出 FontProperties."""
    copied = font.copy()
    copied.set_size(size)
    return copied


def configure_matplotlib():
    """设置 Matplotlib 的基础出图风格."""
    plt.rcParams.update(
        {
            "axes.unicode_minus": False,
            "font.family": "Times New Roman",
            "mathtext.fontset": "stix",
            "axes.linewidth": 0.55,
            "xtick.major.width": 0.5,
            "ytick.major.width": 0.5,
            "hatch.linewidth": 0.45,
            "savefig.dpi": 600,
        }
    )


def read_indices(path):
    """读取 Morris 指标 CSV, 校验必须字段."""
    if not path.is_file():
        raise FileNotFoundError(f"Missing Morris indices file: {path}")
    frame = pd.read_csv(path)
    required_columns = {"metric", "parameter", "mu", "mu_star", "sigma", "mu_star_conf"}
    missing_columns = sorted(required_columns - set(frame.columns))
    if missing_columns:
        joined = ", ".join(missing_columns)
        raise ValueError(f"Missing columns in {path}: {joined}")
    for column in ["mu", "mu_star", "sigma", "mu_star_conf"]:
        frame[column] = pd.to_numeric(frame[column], errors="raise")
    frame["category"] = frame["parameter"].map(PARAMETER_CATEGORY).fillna("Other")
    return frame


def ordered_values(values, preferred_order):
    """按推荐顺序排列标签, 未知标签追加到末尾."""
    unique_values = list(dict.fromkeys(values))
    ordered = [value for value in preferred_order if value in unique_values]
    ordered.extend(value for value in unique_values if value not in ordered)
    return ordered


def metric_label(metric):
    """返回指标的论文图标签."""
    return METRIC_LABELS.get(metric, metric)


def category_colors():
    """返回参数类别配色."""
    palette = sns.color_palette("muted", n_colors=len(CATEGORY_ORDER))
    return {category: palette[index] for index, category in enumerate(CATEGORY_ORDER)}


def category_hatches():
    """返回参数类别填充纹理."""
    return {
        "Crop phenology and leaf": "",
        "Root distribution": "..",
        "Soil hydraulic": "//",
    }


def category_markers():
    """返回参数类别散点标记."""
    return {
        "Crop phenology and leaf": "o",
        "Root distribution": "s",
        "Soil hydraulic": "^",
    }


def parameter_colors(parameters, colors):
    """按参数类别生成颜色列表."""
    result = []
    for parameter in parameters:
        category = PARAMETER_CATEGORY.get(parameter, "Other")
        result.append(colors.get(category, "#777777"))
    return result


def apply_bar_hatches(bars, parameters):
    """给条形图按参数类别添加纹理编码."""
    hatches = category_hatches()
    for patch, parameter in zip(bars.patches, parameters):
        category = PARAMETER_CATEGORY.get(parameter, "Other")
        patch.set_hatch(hatches.get(category, ""))


def category_legend_handles(colors, handle_style="bar"):
    """生成参数类别图例句柄."""
    if handle_style == "marker":
        markers = category_markers()
        return [
            Line2D(
                [0],
                [0],
                marker=markers[category],
                linestyle="",
                markersize=5.0,
                markerfacecolor=colors[category],
                markeredgecolor=BAR_EDGE_COLOR,
                markeredgewidth=0.45,
                label=category,
            )
            for category in CATEGORY_ORDER
        ]

    hatches = category_hatches()
    return [
        Patch(
            facecolor=colors[category],
            edgecolor=BAR_EDGE_COLOR,
            hatch=hatches[category],
            linewidth=0.35,
            label=category,
        )
        for category in CATEGORY_ORDER
    ]


def scatter_by_category(axis, frame, x_column, y_column, colors):
    """按参数类别绘制散点, 同时用颜色和形状编码类别."""
    markers = category_markers()
    if "category" not in frame.columns:
        frame = frame.copy()
        frame["category"] = frame["parameter"].map(PARAMETER_CATEGORY).fillna("Other")
    for category in CATEGORY_ORDER:
        category_frame = frame.loc[frame["category"] == category]
        if category_frame.empty:
            continue
        axis.scatter(
            category_frame[x_column],
            category_frame[y_column],
            s=18,
            marker=markers[category],
            color=colors[category],
            edgecolor=BAR_EDGE_COLOR,
            linewidth=0.3,
            alpha=0.9,
        )


def style_axis(axis, times_font, grid_axis=None):
    """统一坐标轴样式."""
    axis.tick_params(axis="both", which="major", direction="in", length=2.3, width=0.5, labelsize=TICK_SIZE)
    axis.minorticks_off()
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    for side in ["left", "bottom"]:
        axis.spines[side].set_linewidth(0.55)
        axis.spines[side].set_color(BAR_EDGE_COLOR)
    for label in axis.get_xticklabels() + axis.get_yticklabels():
        label.set_fontproperties(font_with_size(times_font, TICK_SIZE))
    if grid_axis:
        axis.grid(axis=grid_axis, color=GRID_COLOR, linewidth=0.35)
        axis.set_axisbelow(True)


def set_axis_text(axis, times_font, title="", xlabel="", ylabel=""):
    """设置英文标题和坐标轴标签."""
    if title:
        axis.set_title(title, fontproperties=font_with_size(times_font, TITLE_SIZE), pad=3)
    if xlabel:
        axis.set_xlabel(xlabel, fontproperties=font_with_size(times_font, LABEL_SIZE), labelpad=2)
    if ylabel:
        axis.set_ylabel(ylabel, fontproperties=font_with_size(times_font, LABEL_SIZE), labelpad=2)


def add_panel_label(axis, label, times_font):
    """添加分面板字母."""
    axis.text(
        0.0,
        1.03,
        f"({label})",
        transform=axis.transAxes,
        ha="left",
        va="bottom",
        fontproperties=font_with_size(times_font, TITLE_SIZE),
        fontweight="bold",
    )


def add_panel_title(axis, label, title, times_font):
    """添加面板字母和短标题."""
    axis.set_title(
        f"({label}) {title}",
        loc="left",
        fontproperties=font_with_size(times_font, TITLE_SIZE),
        pad=4,
    )


def add_category_legend(figure, colors, times_font, anchor_y, columns=3, handle_style="bar"):
    """添加参数类别图例."""
    handles = category_legend_handles(colors, handle_style=handle_style)
    legend = figure.legend(
        handles=handles,
        loc="upper center",
        bbox_to_anchor=(0.5, anchor_y),
        ncol=columns,
        frameon=False,
        prop=font_with_size(times_font, LEGEND_SIZE),
        handletextpad=0.45,
        columnspacing=1.2,
    )
    for text in legend.get_texts():
        text.set_fontproperties(font_with_size(times_font, LEGEND_SIZE))


def add_category_legend_axis(axis, colors, times_font, handle_style="bar"):
    """在空白子图内添加参数类别图例."""
    axis.axis("off")
    handles = category_legend_handles(colors, handle_style=handle_style)
    legend = axis.legend(
        handles=handles,
        loc="center",
        frameon=False,
        prop=font_with_size(times_font, LEGEND_SIZE),
        handletextpad=0.45,
        labelspacing=0.9,
        borderaxespad=0.0,
    )
    for text in legend.get_texts():
        text.set_fontproperties(font_with_size(times_font, LEGEND_SIZE))


def add_bottom_category_legend(figure, colors, times_font, handle_style="bar"):
    """在图底部添加参数类别图例."""
    handles = category_legend_handles(colors, handle_style=handle_style)
    legend = figure.legend(
        handles=handles,
        loc="lower center",
        bbox_to_anchor=(0.5, 0.015),
        ncol=3,
        frameon=False,
        prop=font_with_size(times_font, LEGEND_SIZE),
        handletextpad=0.45,
        columnspacing=1.1,
    )
    for text in legend.get_texts():
        text.set_fontproperties(font_with_size(times_font, LEGEND_SIZE))


def subplot_grid(metric_count):
    """根据指标数量返回散点类多面板图网格尺寸."""
    if metric_count == 7:
        return 4, 2
    columns = 3
    rows = math.ceil(metric_count / columns)
    return rows, columns


def add_common_axis_labels(figure, xlabel, ylabel, times_font, xlabel_y=0.075, ylabel_x=0.024):
    """为多面板散点图添加公共坐标轴标签."""
    if xlabel:
        figure.text(
            0.535,
            xlabel_y,
            xlabel,
            ha="center",
            va="center",
            fontproperties=font_with_size(times_font, LABEL_SIZE),
        )
    if ylabel:
        figure.text(
            ylabel_x,
            0.545,
            ylabel,
            ha="center",
            va="center",
            rotation="vertical",
            fontproperties=font_with_size(times_font, LABEL_SIZE),
        )


def add_legend_to_empty_axes(flat_axes, used_count, colors, times_font, handle_style="bar"):
    """优先使用空白面板放置图例, 输出是否已放置."""
    empty_axes = flat_axes[used_count:]
    if len(empty_axes) == 0:
        return False
    add_category_legend_axis(empty_axes[0], colors, times_font, handle_style=handle_style)
    for axis in empty_axes[1:]:
        axis.axis("off")
    return True


def add_inside_category_legend(axis, colors, times_font, handle_style="bar"):
    """在单面板图的图内空白区添加类别图例."""
    handles = category_legend_handles(colors, handle_style=handle_style)
    legend = axis.legend(
        handles=handles,
        loc="lower right",
        bbox_to_anchor=(0.985, 0.04),
        frameon=False,
        prop=font_with_size(times_font, LEGEND_SIZE - 1),
        handlelength=0.95,
        handletextpad=0.35,
        labelspacing=0.35,
        borderaxespad=0.0,
    )
    for text in legend.get_texts():
        text.set_fontproperties(font_with_size(times_font, LEGEND_SIZE - 1))


def save_figure(figure, output_file):
    """保存 PNG 图件."""
    figure.savefig(output_file, dpi=600, bbox_inches="tight")


def rounded_x_limit(max_value, upper_cap=None):
    """计算适合论文图的 x 轴上限."""
    if max_value <= 0:
        return 1.0
    limit = math.ceil(max_value * 1.04 / 0.1) * 0.1
    if upper_cap is not None:
        limit = min(limit, upper_cap)
    return max(limit, 0.1)


def top_metric_frame(frame, metric, top_n):
    """提取单个指标的 top 参数并按绘图顺序返回."""
    metric_frame = frame.loc[frame["metric"] == metric].copy()
    metric_frame["parameter_order"] = metric_frame["parameter"].apply(
        lambda value: PARAMETER_ORDER.index(value) if value in PARAMETER_ORDER else len(PARAMETER_ORDER)
    )
    top = metric_frame.sort_values(["mu_star", "parameter_order"], ascending=[False, True]).head(top_n)
    return top.sort_values("mu_star", ascending=True)


def plot_ranking_figure(frame, output_file, times_font, top_n):
    """绘制每个输出指标的 Morris mu_star 排名图."""
    metrics = ordered_values(frame["metric"], METRIC_ORDER)
    colors = category_colors()
    max_with_conf = float((frame["mu_star"] + frame["mu_star_conf"]).max())
    x_limit = rounded_x_limit(max_with_conf, upper_cap=1.0 if max_with_conf <= 1.0 else None)
    rows, columns = subplot_grid(len(metrics))
    figure_width = 7.2 if columns == 2 else 7.0
    figure_height = 1.48 * rows + 0.35
    figure, axes = plt.subplots(rows, columns, figsize=(figure_width, figure_height), sharex=True)
    figure.subplots_adjust(left=0.17, right=0.99, top=0.96, bottom=0.085, wspace=0.42, hspace=0.58)
    flat_axes = axes.ravel() if hasattr(axes, "ravel") else [axes]

    for axis, metric, panel_label in zip(flat_axes, metrics, PANEL_LABELS):
        plot_frame = top_metric_frame(frame, metric, top_n)
        y_positions = np.arange(len(plot_frame))
        parameters = plot_frame["parameter"].tolist()
        values = plot_frame["mu_star"].to_numpy()
        conf = plot_frame["mu_star_conf"].to_numpy()
        bars = axis.barh(
            y_positions,
            values,
            height=0.56,
            color=parameter_colors(parameters, colors),
            edgecolor=BAR_EDGE_COLOR,
            linewidth=0.3,
        )
        apply_bar_hatches(bars, parameters)
        axis.errorbar(values, y_positions, xerr=conf, fmt="none", ecolor=BAR_EDGE_COLOR, elinewidth=0.55, capsize=1.7)
        axis.set_yticks(y_positions)
        axis.set_yticklabels(parameters)
        axis.set_xlim(0, x_limit)
        style_axis(axis, times_font, grid_axis="x")
        add_panel_title(axis, panel_label, metric_label(metric), times_font)

    add_legend_to_empty_axes(flat_axes, len(metrics), colors, times_font, handle_style="bar")
    add_common_axis_labels(figure, r"$\mu^*$", "", times_font, xlabel_y=0.038)
    save_figure(figure, output_file)
    plt.close(figure)


def top_ratio_frame(frame, metric, top_n):
    """提取单个指标中 mu_star top 参数并计算 sigma/mu_star."""
    top = top_metric_frame(frame, metric, top_n).copy()
    top["sigma_mu_star_ratio"] = top["sigma"] / top["mu_star"].replace(0, np.nan)
    top = top.dropna(subset=["sigma_mu_star_ratio"])
    return top.sort_values("sigma_mu_star_ratio", ascending=True)


def plot_ratio_figure(frame, output_file, times_font, top_n):
    """绘制关键参数的 sigma/mu_star 比值图, 用于判断作用稳定性."""
    metrics = ordered_values(frame["metric"], METRIC_ORDER)
    colors = category_colors()
    ratio_frames = {metric: top_ratio_frame(frame, metric, top_n) for metric in metrics}
    max_ratio = max(float(ratio_frame["sigma_mu_star_ratio"].max()) for ratio_frame in ratio_frames.values())
    x_limit = rounded_x_limit(max_ratio + 0.12)
    rows, columns = subplot_grid(len(metrics))
    figure_width = 7.2 if columns == 2 else 7.0
    figure_height = 1.48 * rows + 0.35
    figure, axes = plt.subplots(rows, columns, figsize=(figure_width, figure_height), sharex=True)
    figure.subplots_adjust(left=0.17, right=0.99, top=0.96, bottom=0.085, wspace=0.42, hspace=0.58)
    flat_axes = axes.ravel() if hasattr(axes, "ravel") else [axes]

    for axis, metric, panel_label in zip(flat_axes, metrics, PANEL_LABELS):
        plot_frame = ratio_frames[metric]
        y_positions = np.arange(len(plot_frame))
        parameters = plot_frame["parameter"].tolist()
        values = plot_frame["sigma_mu_star_ratio"].to_numpy()
        bars = axis.barh(
            y_positions,
            values,
            height=0.56,
            color=parameter_colors(parameters, colors),
            edgecolor=BAR_EDGE_COLOR,
            linewidth=0.3,
        )
        apply_bar_hatches(bars, parameters)
        for reference in [0.5, 1.0]:
            if reference < x_limit:
                axis.axvline(reference, color="#9a9a9a", linestyle="--", linewidth=0.45, zorder=0)
        axis.set_yticks(y_positions)
        axis.set_yticklabels(parameters)
        axis.set_xlim(0, x_limit)
        style_axis(axis, times_font, grid_axis="x")
        add_panel_title(axis, panel_label, metric_label(metric), times_font)
        for y_position, value in zip(y_positions, values):
            axis.text(
                value + x_limit * 0.012,
                y_position,
                f"{value:.2f}",
                va="center",
                ha="left",
                fontproperties=font_with_size(times_font, ANNOTATION_SIZE - 0.3),
            )

    add_legend_to_empty_axes(flat_axes, len(metrics), colors, times_font, handle_style="bar")
    add_common_axis_labels(figure, r"$R=\sigma/\mu^*$", "", times_font, xlabel_y=0.038)
    save_figure(figure, output_file)
    plt.close(figure)


def annotate_top_points(axis, frame, x_column, y_column, times_font, limit):
    """标注当前面板中最重要的参数."""
    label_frame = frame.assign(max_value=frame[[x_column, y_column]].max(axis=1)).sort_values("max_value", ascending=False).head(limit)
    x_upper = axis.get_xlim()[1]
    y_upper = axis.get_ylim()[1]
    used_offsets = [(4, 4), (4, -8), (-4, 4), (-4, -8)]
    for label_index, (_, row) in enumerate(label_frame.iterrows()):
        dx, dy = used_offsets[label_index % len(used_offsets)]
        ha = "left"
        va = "bottom" if dy >= 0 else "top"
        if row[x_column] > x_upper * 0.78:
            dx = -4
            ha = "right"
        if row[y_column] > y_upper * 0.82:
            dy = -8
            va = "top"
        axis.annotate(
            row["parameter"],
            (row[x_column], row[y_column]),
            xytext=(dx, dy),
            textcoords="offset points",
            ha=ha,
            va=va,
            fontproperties=font_with_size(times_font, ANNOTATION_SIZE),
        )


def plot_morris_plane_figure(frame, output_file, times_font):
    """绘制 Morris 平面, 用于识别非线性和交互影响."""
    metrics = ordered_values(frame["metric"], METRIC_ORDER)
    colors = category_colors()
    x_limit = float(frame["mu_star"].max()) * 1.10
    y_limit = float(frame["sigma"].max()) * 1.12
    rows, columns = subplot_grid(len(metrics))
    has_empty_axis = rows * columns > len(metrics)
    figure_width = 6.8 if columns == 2 else 7.4
    bottom_margin = 0.105 if has_empty_axis else 0.16
    figure, axes = plt.subplots(rows, columns, figsize=(figure_width, 2.02 * rows + 0.55))
    figure.subplots_adjust(left=0.145, right=0.985, top=0.95, bottom=bottom_margin, wspace=0.32, hspace=0.52)
    flat_axes = axes.ravel()

    for axis_index, (axis, metric, panel_label) in enumerate(zip(flat_axes, metrics, PANEL_LABELS)):
        metric_frame = frame.loc[frame["metric"] == metric].copy()
        scatter_by_category(axis, metric_frame, "mu_star", "sigma", colors)
        axis.set_xlim(0, x_limit)
        axis.set_ylim(0, y_limit)
        set_axis_text(axis, times_font)
        style_axis(axis, times_font, grid_axis="both")
        add_panel_title(axis, panel_label, metric_label(metric), times_font)

    legend_in_axis = add_legend_to_empty_axes(flat_axes, len(metrics), colors, times_font, handle_style="marker")
    add_common_axis_labels(figure, r"$\mu^*$", r"$\sigma$", times_font, xlabel_y=0.052, ylabel_x=0.092)
    if not legend_in_axis:
        add_bottom_category_legend(figure, colors, times_font, handle_style="marker")
    save_figure(figure, output_file)
    plt.close(figure)


def normalized_mean_sensitivity(frame):
    """计算跨指标归一化平均敏感性.

    输入为 Morris 指标表. 先在每个输出指标内用最大 mu_star 归一化, 再按参数
    汇总均值, 中位数, 最大值和进入前三名的次数.
    """
    work_frame = frame.copy()
    metric_max = work_frame.groupby("metric")["mu_star"].transform("max")
    work_frame["normalized_mu_star"] = work_frame["mu_star"] / metric_max
    ranks = work_frame.copy()
    ranks["rank"] = ranks.groupby("metric")["mu_star"].rank(method="min", ascending=False)
    top_counts = ranks.loc[ranks["rank"] <= 3].groupby("parameter").size().rename("top3_count")
    summary = (
        work_frame.groupby("parameter")
        .agg(
            mean_norm=("normalized_mu_star", "mean"),
            median_norm=("normalized_mu_star", "median"),
            max_norm=("normalized_mu_star", "max"),
        )
        .reset_index()
    )
    summary = summary.merge(top_counts, on="parameter", how="left").fillna({"top3_count": 0})
    summary["top3_count"] = summary["top3_count"].astype(int)
    summary["category"] = summary["parameter"].map(PARAMETER_CATEGORY).fillna("Other")
    return summary.sort_values(["mean_norm", "top3_count", "max_norm"], ascending=[True, True, True])


def ordered_index(value, preferred_order):
    """返回推荐顺序索引, 未知值排在末尾."""
    if value in preferred_order:
        return preferred_order.index(value)
    return len(preferred_order)


def morris_indices_table(frame):
    """整理完整 Morris 指标表.

    输入为去慢轨迹后的 Morris 指标表. 输出包含全部参数和输出指标的 mu,
    mu_star, sigma 与 mu_star_conf, 供补充表直接审阅.
    """
    table = frame.copy()
    table["metric_label"] = table["metric"].map(METRIC_LABELS).fillna(table["metric"])
    table["metric_order"] = table["metric"].map(lambda value: ordered_index(value, METRIC_ORDER))
    table["parameter_order"] = table["parameter"].map(lambda value: ordered_index(value, PARAMETER_ORDER))
    table = table.sort_values(["metric_order", "mu_star"], ascending=[True, False])
    return table[
        [
            "metric_order",
            "metric",
            "metric_label",
            "parameter_order",
            "parameter",
            "category",
            "mu",
            "mu_star",
            "sigma",
            "mu_star_conf",
        ]
    ]


def normalized_metric_table(frame):
    """计算 figure5 的逐输出归一化过程表.

    输入为去慢轨迹后的 Morris 指标表. 输出包含每个输出指标内的最大 mu_star,
    归一化 mu_star, 以及按 mu_star 降序得到的排名.
    """
    table = frame.copy()
    table["metric_label"] = table["metric"].map(METRIC_LABELS).fillna(table["metric"])
    table["metric_order"] = table["metric"].map(lambda value: ordered_index(value, METRIC_ORDER))
    table["parameter_order"] = table["parameter"].map(lambda value: ordered_index(value, PARAMETER_ORDER))
    table["metric_max_mu_star"] = table.groupby("metric")["mu_star"].transform("max")
    table["normalized_mu_star"] = table["mu_star"] / table["metric_max_mu_star"]
    table["rank_within_metric"] = table.groupby("metric")["mu_star"].rank(method="min", ascending=False).astype(int)
    table = table.sort_values(["metric_order", "rank_within_metric", "parameter_order"])
    return table[
        [
            "metric_order",
            "metric",
            "metric_label",
            "rank_within_metric",
            "parameter_order",
            "parameter",
            "category",
            "mu_star",
            "metric_max_mu_star",
            "normalized_mu_star",
        ]
    ]


def normalized_summary_table(frame):
    """整理 figure5 的参数综合敏感性汇总表.

    输入为去慢轨迹后的 Morris 指标表. 输出为每个参数跨输出指标的平均,
    中位数和最大归一化 mu_star, 以及进入前三名的次数.
    """
    table = normalized_mean_sensitivity(frame).copy()
    table["parameter_order"] = table["parameter"].map(lambda value: ordered_index(value, PARAMETER_ORDER))
    table = table.sort_values(["mean_norm", "top3_count", "max_norm"], ascending=[False, False, False])
    return table[
        [
            "parameter_order",
            "parameter",
            "category",
            "mean_norm",
            "median_norm",
            "max_norm",
            "top3_count",
        ]
    ]


def method_notes_table(top_n):
    """生成 Excel 方法说明表.

    输出说明补充表字段含义, 边界条件为 top_n 必须为正整数.
    """
    return pd.DataFrame(
        [
            {
                "item": "Source data",
                "description": "Filtered Morris indices after excluding slow trajectories.",
            },
            {
                "item": "Morris metrics",
                "description": "mu, mu_star, sigma, and mu_star_conf are exported for all 13 parameters and all output metrics.",
            },
            {
                "item": "mu_star_conf",
                "description": "Half-width of the bootstrap confidence interval for mu_star. Current analysis used 1000 resamples and 0.95 confidence level.",
            },
            {
                "item": "Figure1",
                "description": "The Morris mu_star-sigma plane is exported to assess parameter importance together with nonlinearity or interaction effects.",
            },
            {
                "item": "Figure2",
                "description": f"The plotted ranking figure shows the top {top_n} parameters for each output metric by mu_star.",
            },
            {
                "item": "Figure5",
                "description": "Mean normalized sensitivity is derived by dividing mu_star by the maximum mu_star within each output metric, then averaging by parameter.",
            },
            {
                "item": "Figure6",
                "description": f"The sigma/mu_star ratio is plotted for the top {top_n} parameters of each output metric to indicate effect stability among important parameters.",
            },
            {
                "item": "Sigma",
                "description": "sigma is retained in the complete Morris table for assessing nonlinearity or interaction effects.",
            },
        ]
    )


def format_excel_workbook(writer):
    """设置 Excel 工作簿基础格式.

    输入为 pandas ExcelWriter. 输出直接修改工作簿样式, 无返回值.
    """
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=False)
    for sheet in writer.book.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = body_alignment
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"
        for column_cells in sheet.columns:
            header = str(column_cells[0].value)
            values = [str(cell.value) for cell in column_cells[:100] if cell.value is not None]
            max_length = max([len(header), *(len(value) for value in values)])
            width = min(max(max_length + 2, 10), 42)
            sheet.column_dimensions[column_cells[0].column_letter].width = width


def write_supplementary_workbook(frame, output_file, top_n):
    """导出 Morris 敏感性分析补充表.

    输入为去慢轨迹后的 Morris 指标表, 输出为 xlsx 文件. 当目标目录不存在时
    自动创建. 文件包含完整 Morris 指标, figure5 归一化过程, figure5 汇总,
    以及方法说明.
    """
    output_file.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        morris_indices_table(frame).to_excel(writer, sheet_name="Morris_indices", index=False)
        normalized_metric_table(frame).to_excel(writer, sheet_name="Normalized_by_output", index=False)
        normalized_summary_table(frame).to_excel(writer, sheet_name="Normalized_summary", index=False)
        method_notes_table(top_n).to_excel(writer, sheet_name="Method_notes", index=False)
        format_excel_workbook(writer)


def plot_normalized_mean_figure(frame, output_file, times_font):
    """绘制归一化平均敏感性参数排序图."""
    plot_frame = normalized_mean_sensitivity(frame)
    colors = category_colors()
    parameters = plot_frame["parameter"].tolist()
    values = plot_frame["mean_norm"].to_numpy()
    y_positions = np.arange(len(plot_frame))
    x_limit = rounded_x_limit(float(values.max()) + 0.055, upper_cap=1.0)

    figure, axis = plt.subplots(figsize=(5.5, 4.45))
    figure.subplots_adjust(left=0.29, right=0.985, top=0.965, bottom=0.13)
    bars = axis.barh(
        y_positions,
        values,
        height=0.54,
        color=parameter_colors(parameters, colors),
        edgecolor=BAR_EDGE_COLOR,
        linewidth=0.3,
    )
    apply_bar_hatches(bars, parameters)
    axis.set_yticks(y_positions)
    axis.set_yticklabels(parameters)
    axis.set_xlim(0, x_limit)
    axis.set_xlabel(r"Mean normalized $\mu^*$", fontproperties=font_with_size(times_font, LABEL_SIZE), labelpad=4)
    style_axis(axis, times_font, grid_axis="x")

    for y_position, (_, row) in zip(y_positions, plot_frame.iterrows()):
        axis.text(
            row["mean_norm"] + 0.015,
            y_position,
            f"{row['mean_norm']:.2f}",
            va="center",
            ha="left",
            fontproperties=font_with_size(times_font, ANNOTATION_SIZE),
        )

    add_inside_category_legend(axis, colors, times_font, handle_style="bar")
    save_figure(figure, output_file)
    plt.close(figure)


def main():
    """执行 Morris 敏感性分析绘图."""
    args = parse_arguments()
    if args.top_n <= 0:
        raise ValueError("--top-n must be positive.")

    analysis_dir = resolve_path(args.analysis_dir, ROOT_DIR)
    output_dir = resolve_path(args.output_dir, ROOT_DIR) if args.output_dir else analysis_dir / "figures"
    xlsx_file = resolve_path(args.xlsx_file, output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    find_font(SIMSUN_CANDIDATES, "SimSun")
    times_font = find_font(TIMES_CANDIDATES, "Times New Roman")
    configure_matplotlib()

    filtered_frame = read_indices(analysis_dir / args.filtered_file)
    figure_rows = [
        {
            "file": "figure1_mu_star_sigma_plane_without_slow_trajectories.png",
            "description": "Morris mu_star-sigma plane after dropping slow trajectories.",
        },
        {
            "file": "figure2_mu_star_ranking_without_slow_trajectories.png",
            "description": "Top Morris mu_star ranking after dropping slow trajectories.",
        },
        {
            "file": "figure5_normalized_mean_sensitivity_without_slow.png",
            "description": "Mean normalized Morris mu_star ranking after dropping slow trajectories.",
        },
        {
            "file": "figure6_sigma_mu_star_ratio_without_slow_trajectories.png",
            "description": "Sigma to mu_star ratio ranking for top Morris parameters after dropping slow trajectories.",
        },
    ]

    plot_morris_plane_figure(
        filtered_frame,
        output_dir / figure_rows[0]["file"],
        times_font,
    )
    plot_ranking_figure(
        filtered_frame,
        output_dir / figure_rows[1]["file"],
        times_font,
        args.top_n,
    )
    plot_normalized_mean_figure(
        filtered_frame,
        output_dir / figure_rows[2]["file"],
        times_font,
    )
    plot_ratio_figure(
        filtered_frame,
        output_dir / figure_rows[3]["file"],
        times_font,
        args.top_n,
    )
    write_supplementary_workbook(filtered_frame, xlsx_file, args.top_n)

    for row in figure_rows:
        print(f"绘图完成: {output_dir / row['file']}")
    print(f"补充表完成: {xlsx_file}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
