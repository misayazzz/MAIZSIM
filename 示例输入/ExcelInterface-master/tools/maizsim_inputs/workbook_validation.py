"""Excel 参数表 dry-run 校验."""

import re

from openpyxl import load_workbook

from .drip import validate_drip_records_for_run
from .errors import ConfigError
from .validation_rules import REQUIRED_SHEET_FIELDS


def normalize_key(value):
    """把 Excel 表头或 ID 标准化为可比较字符串."""
    return re.sub(r"[^a-z0-9]", "", str(value).strip().lower())


def value_text(value):
    """把单元格值转成去空白字符串."""
    if value is None:
        return ""
    return str(value).strip()


def find_sheet(workbook, sheet_name):
    """按大小写不敏感方式查找工作表."""
    target = normalize_key(sheet_name)
    for existing_name in workbook.sheetnames:
        if normalize_key(existing_name) == target:
            return workbook[existing_name]
    raise ConfigError(f"参数表缺少工作表: {sheet_name}")


def read_table(worksheet):
    """读取工作表表头和记录, 第一行作为表头."""
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = next(rows)
    except StopIteration as exc:
        raise ConfigError(f"{worksheet.title} 工作表为空.") from exc

    normalized_headers = [normalize_key(header) if header is not None else "" for header in headers]
    if not any(normalized_headers):
        raise ConfigError(f"{worksheet.title} 第一行没有有效表头.")

    seen_headers = {}
    for column_number, header in enumerate(normalized_headers, 1):
        if not header:
            continue
        if header in seen_headers:
            raise ConfigError(
                f"{worksheet.title} 存在重复表头 {header}, 列 {seen_headers[header]} 和列 {column_number}."
            )
        seen_headers[header] = column_number

    records = []
    for row_number, row in enumerate(rows, 2):
        record = {"__row_number__": row_number}
        has_value = False
        for header, value in zip(normalized_headers, row):
            if not header:
                continue
            record[header] = value
            if value_text(value):
                has_value = True
        if has_value:
            records.append(record)
    return {"headers": set(seen_headers), "records": records}


def get_records(table):
    """返回表格记录列表."""
    return table["records"]


def require_fields(table, field_names, sheet_name):
    """确认工作表包含指定字段."""
    missing = []
    for field_name in field_names:
        key = normalize_key(field_name)
        if key not in table["headers"]:
            missing.append(field_name)
    if missing:
        joined = ", ".join(missing)
        raise ConfigError(f"{sheet_name} 缺少字段: {joined}")


def get_record_value(record, field_names, context):
    """从记录中读取字段值."""
    for field_name in field_names:
        key = normalize_key(field_name)
        if key in record:
            return value_text(record[key])
    joined = ", ".join(field_names)
    raise ConfigError(f"{context} 缺少字段: {joined}")


def build_index(records, field_names, sheet_name):
    """按指定字段构建值索引."""
    index = {}
    for record in records:
        key = get_record_value(record, field_names, sheet_name)
        if not key:
            raise ConfigError(f"{sheet_name} 第 {record['__row_number__']} 行索引字段为空.")
        if key in index:
            first_row = index[key]["__row_number__"]
            current_row = record["__row_number__"]
            raise ConfigError(f"{sheet_name} 存在重复 ID {key}, 行 {first_row} 和行 {current_row}.")
        index[key] = record
    return index


def build_group_index(records, field_names):
    """按指定字段构建一对多索引."""
    index = {}
    for record in records:
        key = get_record_value(record, field_names, "group index")
        if key:
            index.setdefault(key, []).append(record)
    return index


def require_lookup(value, index, source_context, target_sheet):
    """确认引用值存在于目标索引."""
    if not value:
        raise ConfigError(f"{source_context} 引用值为空.")
    if value not in index:
        raise ConfigError(f"{source_context} 引用 {value} 不存在于 {target_sheet}.")


def require_group_lookup(value, index, source_context, target_sheet):
    """确认引用值存在于一对多索引."""
    if not value:
        raise ConfigError(f"{source_context} 引用值为空.")
    if value not in index or not index[value]:
        raise ConfigError(f"{source_context} 引用 {value} 不存在于 {target_sheet}.")


def require_climate_location(climate_records, climate_id, location, context):
    """确认 Climate 中存在指定 ClimateID 和 Location 的组合."""
    for record in climate_records:
        record_climate_id = get_record_value(record, ["ClimateID"], "Climate")
        record_location = get_record_value(record, ["Location"], "Climate")
        if record_climate_id == climate_id and record_location == location:
            return
    raise ConfigError(f"{context} 在 Climate 中找不到 ClimateID={climate_id}, Location={location} 的记录.")


def require_weather_record(weather_records, climate_id, weather_id, context):
    """确认 Weather 中存在指定 ClimateID 和 WeatherID 的组合."""
    for record in weather_records:
        record_climate_id = get_record_value(record, ["ClimateID"], "Weather")
        record_weather_id = get_record_value(record, ["WeatherID"], "Weather")
        if record_climate_id == climate_id and record_weather_id == weather_id:
            return record
    raise ConfigError(f"{context} 在 Weather 中找不到 ClimateID={climate_id}, WeatherID={weather_id} 的记录.")


def validate_workbook(input_excel_file, selected_ids):
    """校验参数表关键 sheet, selected IDs 和跨表引用."""
    workbook = load_workbook(input_excel_file, read_only=True, data_only=True)
    try:
        tables_by_sheet = {}
        for sheet_name, required_fields in REQUIRED_SHEET_FIELDS.items():
            worksheet = find_sheet(workbook, sheet_name)
            table = read_table(worksheet)
            require_fields(table, required_fields, sheet_name)
            tables_by_sheet[sheet_name] = table

        description_index = build_index(get_records(tables_by_sheet["Description"]), ["ID"], "Description")
        climate_index = build_index(get_records(tables_by_sheet["Climate"]), ["ClimateID"], "Climate")
        soil_groups = build_group_index(get_records(tables_by_sheet["Soil"]), ["SoilFile"])
        grid_index = build_index(get_records(tables_by_sheet["GridRatio"]), ["SoilFile"], "GridRatio")
        variety_index = build_index(get_records(tables_by_sheet["Variety"]), ["Hybrid"], "Variety")
        biology_index = build_index(get_records(tables_by_sheet["Biology"]), ["ID"], "Biology")
        solute_index = build_index(get_records(tables_by_sheet["Solute"]), ["ID"], "Solute")
        gas_index = build_index(get_records(tables_by_sheet["Gas"]), ["ID"], "Gas")
        mulch_geo_index = build_index(get_records(tables_by_sheet["MulchGeo"]), ["ID"], "MulchGeo")
        mulch_decomp_index = build_index(get_records(tables_by_sheet["MulchDecomp"]), ["ID"], "MulchDecomp")
        water_mov_index = build_index(get_records(tables_by_sheet["WaterMovParam"]), ["ID"], "WaterMovParam")
        tillage_index = build_index(get_records(tables_by_sheet["Tillage"]), ["ID"], "Tillage")
        time_index = build_index(get_records(tables_by_sheet["Time"]), ["ID"], "Time")
        init_index = build_index(get_records(tables_by_sheet["Init"]), ["ID"], "Init")
        fertilization_groups = build_group_index(get_records(tables_by_sheet["Fertilization"]), ["ID"])
        drip_groups = build_group_index(get_records(tables_by_sheet["Drip"]), ["ID"])
        drip_node_groups = build_group_index(get_records(tables_by_sheet["DripNodes"]), ["ID"])
        weather_records = get_records(tables_by_sheet["Weather"])
        climate_records = get_records(tables_by_sheet["Climate"])

        runs = []
        for selected_id in selected_ids:
            if selected_id not in description_index:
                raise ConfigError(f"selected_id {selected_id} 不存在于 Description.ID.")

            record = description_index[selected_id]
            context = f"Description.ID={selected_id}"
            weather_id = get_record_value(record, ["WeatherID"], context)
            climate_id = get_record_value(record, ["ClimateID"], context)
            soil_file = get_record_value(record, ["SoilFile"], context)
            hybrid = get_record_value(record, ["Hybrid"], context)
            biology = get_record_value(record, ["Biology"], context)
            solute = get_record_value(record, ["Solute"], context)
            gas_co2 = get_record_value(record, ["Gas_CO2"], context)
            gas_o2 = get_record_value(record, ["Gas_O2"], context)
            gas_n2o = get_record_value(record, ["Gas_N2O"], context)
            mulch_geo = get_record_value(record, ["MulchGeo"], context)
            mulch_decomp = get_record_value(record, ["MulchDecomp"], context)
            water_mov_param = get_record_value(record, ["WaterMovParam"], context)
            tillage = get_record_value(record, ["Tillage"], context)
            location = get_record_value(record, ["Location"], context)
            run_path = get_record_value(record, ["Path"], context)

            for field_name in ["WeatherFileName", "ClimateFile", "NitrogenFile", "VarietyFile", "SoilName", "Gas_File"]:
                get_record_value(record, [field_name], context)

            require_lookup(climate_id, climate_index, f"{context}.ClimateID", "Climate.ClimateID")
            weather_record = require_weather_record(weather_records, climate_id, weather_id, context)
            climate_record = climate_index[climate_id]
            time_record = time_index[selected_id]
            require_group_lookup(soil_file, soil_groups, f"{context}.SoilFile", "Soil.SoilFile")
            require_lookup(soil_file, grid_index, f"{context}.SoilFile", "GridRatio.SoilFile")
            require_lookup(hybrid, variety_index, f"{context}.Hybrid", "Variety.Hybrid")
            require_lookup(biology, biology_index, f"{context}.Biology", "Biology.ID")
            require_lookup(solute, solute_index, f"{context}.Solute", "Solute.ID")
            require_lookup(gas_co2, gas_index, f"{context}.Gas_CO2", "Gas.ID")
            require_lookup(gas_o2, gas_index, f"{context}.Gas_O2", "Gas.ID")
            require_lookup(gas_n2o, gas_index, f"{context}.Gas_N2O", "Gas.ID")
            require_lookup(mulch_geo, mulch_geo_index, f"{context}.MulchGeo", "MulchGeo.ID")
            require_lookup(mulch_decomp, mulch_decomp_index, f"{context}.MulchDecomp", "MulchDecomp.ID")
            require_lookup(water_mov_param, water_mov_index, f"{context}.WaterMovParam", "WaterMovParam.ID")
            require_lookup(tillage, tillage_index, f"{context}.Tillage", "Tillage.ID")
            require_lookup(selected_id, time_index, f"{context}.ID", "Time.ID")
            require_lookup(selected_id, init_index, f"{context}.ID", "Init.ID")
            require_group_lookup(selected_id, fertilization_groups, f"{context}.ID", "Fertilization.ID")
            require_climate_location(climate_records, climate_id, location, context)

            drip_records = drip_groups.get(selected_id, [])
            drip_node_records = drip_node_groups.get(selected_id, [])
            validate_drip_records_for_run(selected_id, drip_records, drip_node_records)

            weather_climate_id = get_record_value(weather_record, ["ClimateID"], f"Weather.WeatherID={weather_id}")
            if weather_climate_id and weather_climate_id != climate_id:
                raise ConfigError(
                    f"{context} 的 ClimateID={climate_id}, 但 Weather.WeatherID={weather_id} 对应 ClimateID={weather_climate_id}."
                )

            runs.append(
                {
                    "id": selected_id,
                    "path": run_path,
                    "soil_file": soil_file,
                    "weather_file_name": get_record_value(record, ["WeatherFileName"], context),
                    "weather_source_name": get_record_value(
                        weather_record, ["Source_name"], f"Weather.WeatherID={weather_id}"
                    ),
                    "weather_time": get_record_value(weather_record, ["Time"], f"Weather.WeatherID={weather_id}"),
                    "climate_id": climate_id,
                    "weather_id": weather_id,
                    "start_date": time_record[normalize_key("startDate")],
                    "end_date": time_record[normalize_key("EndDate")],
                    "daily_wind": climate_record[normalize_key("DailyWind")],
                    "rel_humid": climate_record[normalize_key("RelHumid")],
                    "daily_co2": climate_record[normalize_key("DailyCO2")],
                    "drip_records": drip_records,
                    "drip_node_records": drip_node_records,
                }
            )

        return {
            "selected_count": len(selected_ids),
            "available_run_count": len(description_index),
            "runs": runs,
        }
    finally:
        workbook.close()
