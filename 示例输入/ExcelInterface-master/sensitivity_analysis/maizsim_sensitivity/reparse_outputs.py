"""从已有样本输出重解析扩展指标并重新计算 Morris 指数."""

from concurrent.futures import ThreadPoolExecutor
from concurrent.futures import as_completed
from datetime import datetime
from pathlib import Path
import argparse
import math
import sys
import time

from SALib.analyze import morris
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
import pandas as pd

from .config import SensitivityConfigError
from .config import read_config
from .metric_catalog import metric_group
from .metric_catalog import metric_label
from .metric_catalog import metric_metadata_rows
from .metric_catalog import morris_metric_names
from .outputs import parse_metrics
from .runner import default_max_workers
from .sampling import analysis_matrix_from_sample_table
from .sampling import build_problem


EXTENDED_OUTPUTS_FILE = "outputs_extended.csv"
EXTENDED_ALL_FILE = "morris_indices_extended_all.csv"
EXTENDED_FILTERED_FILE = "morris_indices_extended_without_slow_trajectories.csv"
EXTENDED_SUMMARY_FILE = "morris_indices_extended_generation_summary.csv"
MISSING_SUMMARY_FILE = "metric_missing_summary.csv"
FAILURES_FILE = "outputs_extended_failures.csv"
WORKBOOK_FILE = "morris_sensitivity_extended_tables.xlsx"
FAILURE_COLUMNS = ["sample_id", "sample_dir", "started_at", "finished_at", "elapsed_seconds", "status", "error"]


def parse_args(arguments=None):
    """解析重解析命令行参数."""
    parser = argparse.ArgumentParser(
        description="Reparse existing MAIZSIM sample outputs and calculate extended Morris indices."
    )
    parser.add_argument("--bounds", required=True, help="Path to sensitivity_parameter_bounds.toml.")
    parser.add_argument("--analysis-dir", required=True, help="Existing sensitivity analysis directory.")
    parser.add_argument("--max-workers", type=int, help="Parallel parser worker count.")
    parser.add_argument("--limit-samples", type=int, help="Debug only: parse the first N samples.")
    parser.add_argument(
        "--skip-workbook",
        action="store_true",
        help="Do not write the extended Excel workbook.",
    )
    return parser.parse_args(arguments)


def write_csv(path, rows_or_frame):
    """写入 UTF-8 BOM CSV."""
    frame = rows_or_frame if isinstance(rows_or_frame, pd.DataFrame) else pd.DataFrame(rows_or_frame)
    frame.to_csv(path, index=False, encoding="utf-8-sig")


def write_failure_csv(path, rows):
    """写出解析失败样本表, 即使没有失败也保留表头."""
    frame = pd.DataFrame(rows, columns=FAILURE_COLUMNS)
    frame.to_csv(path, index=False, encoding="utf-8-sig")


def parser_worker_count(args):
    """确定解析 worker 数."""
    if args.max_workers is not None:
        if args.max_workers <= 0:
            raise SensitivityConfigError("--max-workers 必须是正整数.")
        return args.max_workers
    return min(4, default_max_workers())


def load_sample_table(analysis_dir, limit_samples=None):
    """读取样本表."""
    path = analysis_dir / "samples.csv"
    if not path.is_file():
        raise SensitivityConfigError(f"samples.csv 不存在: {path}")
    frame = pd.read_csv(path)
    if limit_samples is not None:
        if limit_samples <= 0:
            raise SensitivityConfigError("--limit-samples 必须是正整数.")
        frame = frame.head(limit_samples).copy()
    return frame


def parse_one_sample(analysis_dir, sample_id):
    """解析一个已有样本目录."""
    sample_dir = analysis_dir / "samples" / f"sample_{int(sample_id):06d}"
    started_at = datetime.now()
    start = time.perf_counter()
    result = {
        "sample_id": int(sample_id),
        "sample_dir": str(sample_dir),
        "started_at": started_at.isoformat(timespec="seconds"),
        "finished_at": "",
        "elapsed_seconds": "",
        "status": "failed",
        "error": "",
    }
    try:
        if not sample_dir.is_dir():
            raise FileNotFoundError(f"样本目录不存在: {sample_dir}")
        metrics = parse_metrics(sample_dir)
        result["status"] = "success"
        return result, {"sample_id": int(sample_id), **metrics}
    except Exception as exc:
        result["error"] = str(exc)
        return result, None
    finally:
        result["finished_at"] = datetime.now().isoformat(timespec="seconds")
        result["elapsed_seconds"] = f"{time.perf_counter() - start:.3f}"


def reparse_samples(analysis_dir, sample_table, max_workers):
    """并行重解析全部样本输出."""
    manifest_rows = []
    output_rows = []
    sample_ids = sample_table["sample_id"].astype(int).tolist()
    completed = 0
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        future_map = {
            executor.submit(parse_one_sample, analysis_dir, sample_id): sample_id
            for sample_id in sample_ids
        }
        for future in as_completed(future_map):
            manifest, metrics = future.result()
            manifest_rows.append(manifest)
            if metrics is not None:
                output_rows.append(metrics)
            completed += 1
            if completed == len(sample_ids) or completed % 100 == 0:
                print(f"重解析进度: {completed}/{len(sample_ids)}")
    manifest_rows.sort(key=lambda row: row["sample_id"])
    output_rows.sort(key=lambda row: row["sample_id"])
    return manifest_rows, output_rows


def slow_trajectory_ids(analysis_dir):
    """读取慢轨迹编号."""
    path = analysis_dir / "slow_samples.csv"
    if not path.is_file():
        return []
    frame = pd.read_csv(path)
    if "trajectory_id" not in frame.columns:
        return []
    return sorted(frame["trajectory_id"].dropna().astype(int).unique().tolist())


def filter_without_slow(sample_table, output_frame, dropped_trajectory_ids):
    """剔除包含慢样本的整条 Morris 轨迹."""
    if not dropped_trajectory_ids:
        return sample_table.copy(), output_frame.copy()
    filtered_sample_table = sample_table.loc[~sample_table["trajectory_id"].isin(dropped_trajectory_ids)].copy()
    kept_sample_ids = set(filtered_sample_table["sample_id"].astype(int))
    filtered_output_frame = output_frame.loc[output_frame["sample_id"].astype(int).isin(kept_sample_ids)].copy()
    return filtered_sample_table, filtered_output_frame


def value_is_missing(value):
    """判断数值是否缺失."""
    return value is None or (isinstance(value, float) and math.isnan(value))


def zero_indices(problem, metric_name):
    """为零方差指标生成零敏感性结果."""
    return [
        {
            "metric": metric_name,
            "parameter": parameter_name,
            "mu": 0.0,
            "mu_star": 0.0,
            "sigma": 0.0,
            "mu_star_conf": 0.0,
        }
        for parameter_name in problem["names"]
    ]


def calculate_extended_indices(problem, sample_table, output_frame, metric_names, defaults):
    """计算扩展 Morris 指数, 并返回缺失指标摘要."""
    sorted_outputs = output_frame.sort_values("sample_id").copy()
    sorted_samples = sample_table.sort_values("sample_id").copy()
    missing_rows = []
    if len(sorted_outputs) != len(sorted_samples):
        raise SensitivityConfigError(
            f"输出样本数和参数样本数不一致: outputs={len(sorted_outputs)}, samples={len(sorted_samples)}"
        )
    input_matrix = analysis_matrix_from_sample_table(defaults["parameters"], sorted_samples)
    index_rows = []
    for metric_name in metric_names:
        if metric_name not in sorted_outputs.columns:
            missing_rows.append(
                {
                    "metric": metric_name,
                    "reason": "missing_column",
                    "sample_count": len(sorted_outputs),
                    "missing_count": len(sorted_outputs),
                }
            )
            continue
        values = pd.to_numeric(sorted_outputs[metric_name], errors="coerce")
        missing_count = int(values.isna().sum())
        if missing_count:
            missing_rows.append(
                {
                    "metric": metric_name,
                    "reason": "missing_values",
                    "sample_count": len(values),
                    "missing_count": missing_count,
                }
            )
            continue
        value_array = values.to_numpy(dtype=float)
        if len(set(value_array.tolist())) <= 1:
            index_rows.extend(zero_indices(problem, metric_name))
            continue
        analysis = morris.analyze(
            problem,
            input_matrix,
            value_array,
            num_resamples=defaults.get("analysis_num_resamples", 1000),
            conf_level=defaults.get("analysis_conf_level", 0.95),
            scaled=defaults.get("analysis_scaled", False),
            print_to_console=False,
            num_levels=defaults.get("num_levels", 4),
            seed=defaults.get("seed", 20260429),
        )
        for index, parameter_name in enumerate(analysis["names"]):
            index_rows.append(
                {
                    "metric": metric_name,
                    "parameter": parameter_name,
                    "mu": analysis["mu"][index],
                    "mu_star": analysis["mu_star"][index],
                    "sigma": analysis["sigma"][index],
                    "mu_star_conf": analysis["mu_star_conf"][index],
                }
            )
    indices = pd.DataFrame(index_rows)
    if not indices.empty:
        indices = indices.sort_values(["metric", "mu_star"], ascending=[True, False])
    return indices, pd.DataFrame(missing_rows)


def add_index_metadata(frame):
    """添加指标标签、分组和排名列."""
    if frame.empty:
        return frame.copy()
    result = frame.copy()
    result["metric_group"] = result["metric"].map(metric_group)
    result["metric_label"] = result["metric"].map(metric_label)
    result["rank"] = result.groupby("metric")["mu_star"].rank(method="min", ascending=False).astype(int)
    return result.sort_values(["metric_group", "metric", "rank", "parameter"])


def format_excel_workbook(writer):
    """设置 Excel 工作簿基础格式."""
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=False)
    for sheet in writer.book.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = body_alignment
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"
        for column_cells in sheet.columns:
            header = str(column_cells[0].value)
            values = [str(cell.value) for cell in column_cells[:100] if cell.value is not None]
            width = min(max([len(header), *(len(value) for value in values)]) + 2, 46)
            sheet.column_dimensions[column_cells[0].column_letter].width = max(width, 10)


def safe_sheet_name(value, used_names):
    """生成 Excel 安全 sheet 名."""
    base = "".join(char if char not in '[]:*?/\\' else "_" for char in value)[:31]
    name = base or "Sheet"
    suffix = 1
    while name in used_names:
        tail = f"_{suffix}"
        name = f"{base[:31 - len(tail)]}{tail}"
        suffix += 1
    used_names.add(name)
    return name


def write_extended_workbook(filtered_indices, all_indices, missing_summary, output_file):
    """写出扩展敏感性 Excel 表."""
    output_file.parent.mkdir(parents=True, exist_ok=True)
    filtered = add_index_metadata(filtered_indices)
    all_frame = add_index_metadata(all_indices)
    metadata = pd.DataFrame(metric_metadata_rows())
    used_names = set()
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        filtered.to_excel(writer, sheet_name=safe_sheet_name("Filtered_indices", used_names), index=False)
        all_frame.to_excel(writer, sheet_name=safe_sheet_name("All_indices", used_names), index=False)
        metadata.to_excel(writer, sheet_name=safe_sheet_name("Metric_metadata", used_names), index=False)
        missing_summary.to_excel(writer, sheet_name=safe_sheet_name("Missing_summary", used_names), index=False)
        for group_name, group_frame in filtered.groupby("metric_group", sort=True):
            group_frame.to_excel(writer, sheet_name=safe_sheet_name(group_name, used_names), index=False)
        format_excel_workbook(writer)


def write_generation_summary(analysis_dir, all_indices, filtered_indices, sample_table, filtered_sample_table, dropped_ids):
    """写出扩展 Morris 生成摘要."""
    rows = [
        {
            "version": "all_samples",
            "output_file": EXTENDED_ALL_FILE,
            "sample_count": len(sample_table),
            "trajectory_count": sample_table["trajectory_id"].nunique(),
            "dropped_trajectory_ids": "",
            "morris_index_rows": len(all_indices),
        },
        {
            "version": "without_slow_trajectories",
            "output_file": EXTENDED_FILTERED_FILE,
            "sample_count": len(filtered_sample_table),
            "trajectory_count": filtered_sample_table["trajectory_id"].nunique(),
            "dropped_trajectory_ids": ",".join(str(item) for item in dropped_ids),
            "morris_index_rows": len(filtered_indices),
        },
    ]
    write_csv(analysis_dir / EXTENDED_SUMMARY_FILE, rows)


def main(arguments=None):
    """执行扩展输出重解析."""
    args = parse_args(arguments)
    try:
        config = read_config(Path(args.bounds))
        analysis_dir = Path(args.analysis_dir).expanduser().resolve()
        if not analysis_dir.is_dir():
            raise SensitivityConfigError(f"分析目录不存在: {analysis_dir}")
        sample_table = load_sample_table(analysis_dir, args.limit_samples)
        max_workers = parser_worker_count(args)
        print(f"分析目录: {analysis_dir}")
        print(f"样本数: {len(sample_table)}, parser workers={max_workers}")
        manifest_rows, output_rows = reparse_samples(analysis_dir, sample_table, max_workers)
        failures = [row for row in manifest_rows if row["status"] != "success"]
        output_frame = pd.DataFrame(output_rows).sort_values("sample_id")
        write_csv(analysis_dir / EXTENDED_OUTPUTS_FILE, output_frame)
        write_failure_csv(analysis_dir / FAILURES_FILE, failures)
        if failures:
            print(f"存在解析失败样本, 已写入 {FAILURES_FILE}. failed={len(failures)}")
            return 1

        problem = build_problem(config["parameters"])
        defaults = dict(config["morris_defaults"])
        defaults["parameters"] = config["parameters"]
        metric_names = morris_metric_names(include_legacy=True)
        all_indices, all_missing = calculate_extended_indices(
            problem,
            sample_table,
            output_frame,
            metric_names,
            defaults,
        )
        all_indices.to_csv(analysis_dir / EXTENDED_ALL_FILE, index=False, encoding="utf-8-sig")

        dropped_ids = slow_trajectory_ids(analysis_dir)
        filtered_sample_table, filtered_output_frame = filter_without_slow(sample_table, output_frame, dropped_ids)
        filtered_indices, filtered_missing = calculate_extended_indices(
            problem,
            filtered_sample_table,
            filtered_output_frame,
            metric_names,
            defaults,
        )
        filtered_indices.to_csv(analysis_dir / EXTENDED_FILTERED_FILE, index=False, encoding="utf-8-sig")
        missing_summary = pd.concat(
            [
                all_missing.assign(version="all_samples"),
                filtered_missing.assign(version="without_slow_trajectories"),
            ],
            ignore_index=True,
        )
        write_csv(analysis_dir / MISSING_SUMMARY_FILE, missing_summary)
        write_generation_summary(analysis_dir, all_indices, filtered_indices, sample_table, filtered_sample_table, dropped_ids)
        if not args.skip_workbook:
            write_extended_workbook(filtered_indices, all_indices, missing_summary, analysis_dir / WORKBOOK_FILE)

        print(f"扩展 outputs 完成: {analysis_dir / EXTENDED_OUTPUTS_FILE}")
        print(f"扩展 Morris all 完成: {analysis_dir / EXTENDED_ALL_FILE}")
        print(f"扩展 Morris 去慢轨迹完成: {analysis_dir / EXTENDED_FILTERED_FILE}")
        if not args.skip_workbook:
            print(f"扩展表格完成: {analysis_dir / WORKBOOK_FILE}")
        return 0
    except SensitivityConfigError as exc:
        sys.stderr.write(f"错误: {exc}\n")
        return 1
    except Exception as exc:
        sys.stderr.write(f"未处理错误: {exc}\n")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
